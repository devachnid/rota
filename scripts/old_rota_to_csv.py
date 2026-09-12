#!/usr/bin/env python3
"""Step 1 of the one-off move from the old Google Sheets rota: read one
year tab of the .xlsx export and write the cells for a date window as one
long CSV, which `manage.py import_rota` then reads.

The tab's shape: row 1 holds each date on its AM column (merged across AM
and PM), row 2 the weekday, row 3 "AM"/"PM", and from row 4 down one
person per row with the name in column A and a label in each cell. Cell
comments come out as the note column. Blank cells are dropped; "Off" is
kept, so the importer can check that a row it was told to skip really has
nothing in it.

Needs openpyxl, which the app itself does not: run it anywhere with
    pip install openpyxl
    python scripts/old_rota_to_csv.py rota.xlsx "Rota 2026" 2026-09-21 2026-10-23 window.csv
The CSV holds people's names against dates, so keep it out of the repo.
"""

import csv
import datetime as dt
import sys


def main(xlsx, sheet, start, end, out):
    import openpyxl  # only here, so the app never depends on it
    start, end = dt.date.fromisoformat(start), dt.date.fromisoformat(end)
    ws = openpyxl.load_workbook(xlsx, data_only=True)[sheet]
    columns = []  # (date, part, column index)
    for c in range(2, ws.max_column + 1):
        v = ws.cell(1, c).value
        if isinstance(v, dt.datetime) and start <= v.date() <= end:
            if (ws.cell(3, c).value, ws.cell(3, c + 1).value) != ("AM", "PM"):
                sys.exit(f"column {c}: expected AM/PM under {v.date()}")
            columns += [(v.date(), "AM", c), (v.date(), "PM", c + 1)]
    if not columns:
        sys.exit(f"no dates between {start} and {end} in row 1 of {sheet!r}")
    rows = 0
    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["date", "part", "row", "label", "note"])
        for r in range(4, ws.max_row + 1):
            name = " ".join(str(ws.cell(r, 1).value or "").split())
            if not name:
                continue
            for day, part, c in columns:
                cell = ws.cell(r, c)
                label = " ".join(str(cell.value).split()) if cell.value is not None else ""
                if not label:
                    continue
                note = " ".join(cell.comment.text.split()) if cell.comment else ""
                w.writerow([day.isoformat(), part, name, label, note])
                rows += 1
    print(f"{rows} cells over {len(columns) // 2} days for "
          f"{len({r for r in range(4, ws.max_row + 1) if ws.cell(r, 1).value})} rows -> {out}")


if __name__ == "__main__":
    if len(sys.argv) != 6:
        sys.exit(__doc__)
    main(*sys.argv[1:])
